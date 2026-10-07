#!/usr/bin/env python3
"""Somekalenteri marraskuu 2026 – tammikuu 2027 (Excel). python3 kalenteri_build.py <tiedosto.xlsx>"""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

FB_IG = 'Facebook + Instagram'
VIDEO = 'Instagram Reels + TikTok'
TT = 'TikTok'
KAIKKI = 'Facebook + Instagram + TikTok'

H = '#ikkunakauppias #ikkunat #ovet #ikkunaremontti #skaala #turku #varsinaissuomi'

# (pvm, kanavat, muoto, teema, aihe, teksti, kuva/video, linkki/CTA, tehostus €)
P = [
    # --- Vko 45: lanseeraus ---
    ('2026-11-02', FB_IG, 'Kuva', 'Tutuksi', 'Lanseeraus: Ikkunoiden & ovien kauppias',
     'Ikkunoiden & ovien kauppias 🪟\n\nOlemme Nordic Ikkunat & Ovet Oy ja myymme ja asennamme Skaalan ikkunoita ja ovia. Ne valmistetaan mittatilaustyönä Suomessa, ja asennuksen hoitaa ammattilainen alusta loppuun.\n\nKotimaisuus & ammattitaito kulkee käsi kädessä. Seuraa tiliä, niin saat vinkit ikkuna- ja oviremonttiin suoraan syötteeseesi. 👋',
     'ks-postaus.png (somepaketti)', 'Seuraa tiliä', 20),
    ('2026-11-04', VIDEO, 'Video 20–30 s', 'Tutuksi', 'Sami esittäytyy',
     'Moi! Olen Sami Elmeranta ja vastaan Ikkunakauppias.fi:stä. Kerron täällä suoraan, mitä ikkunoiden ja ovien vaihto oikeasti vaatii – ja milloin se ei vielä kannata. Kysy kommenteissa mitä vain! 👇',
     'Sami kameralle, taustalla ikkuna tai auto. Puhu vapaasti, 3 lausetta: kuka, mitä teemme, mitä tältä tililtä saa.', 'Kysy kommenteissa', 0),
    ('2026-11-06', FB_IG, 'Karuselli (4 kuvaa)', 'Tutuksi', 'Näin ikkunaremontti etenee meillä',
     'Näin ikkunaremontti etenee meillä 👉\n\n1. Ilmainen mittauskäynti\n2. Kirjallinen, sitoumukseton tarjous\n3. Ikkunat valmistetaan mittojen mukaan Suomessa\n4. Asennus, tiivistys ja siivous\n\nYksi ikkuna vaihdetaan yleensä muutamassa tunnissa, joten koti ei ehdi jäähtyä.',
     'Kuvat: ig-2-mittauskaynti.png + 3 omaa kuvaa (mittanauha, tarjous, asennus)', 'ikkunakauppias.fi', 0),
    ('2026-11-08', FB_IG, 'Kuva', 'Kausi', 'Isänpäivä',
     'Hyvää isänpäivää kaikille isille – erityisesti niille, jotka tietävät aina, missä ruuvimeisseli on. 🔧💛',
     'Lämmin kuva: työkalupakki ikkunalaudalla tai logo-kuva', '–', 0),
    # --- Vko 46: talvi tulee, veto ja huurre ---
    ('2026-11-09', FB_IG, 'Karuselli (3 kuvaa)', 'Vinkki', 'Missä huurre on? 3 paikkaa, 3 merkitystä',
     'Huurre kertoo ikkunan kunnosta enemmän kuin luulet ❄️\n\n1️⃣ Ulkolasin ulkopinnalla: hyvä merkki, lasi eristää.\n2️⃣ Lasien välissä: tiivisteet tai lasielementti vuotavat.\n3️⃣ Sisäpinnalla: lasi on kylmä tai sisäilma kostea.\n\nMissä teillä huurtuu? Kerro kommenteissa.',
     'Kolme lähikuvaa huurteesta tai yksinkertainen grafiikka (Canva), sävyt somepaketista', 'Kommentoi', 15),
    ('2026-11-11', VIDEO, 'Video 15–25 s', 'Vinkki', 'Vetääkö ikkunasta? Kynttilätesti',
     'Vetääkö ikkunasta? Tee tämä testi ennen kuin vaihdat 🕯️ Jos liekki värisee puitteen reunalla, tiivisteet ovat väsyneet. Ne voi vaihtaa. Jos puite on vääntynyt, vaihto on järkevämpi.',
     'Kansi tiktok-1-veto.png. Kuvaa kynttilä puitteen reunalla.', 'Ilmainen arvio: linkki profiilissa', 0),
    ('2026-11-13', FB_IG, 'Kuva', 'Vinkki', '3 merkkiä, että ikkunat kannattaa uusia',
     'Kannattaako teillä vaihtaa ikkunat? 3 merkkiä:\n\n1️⃣ Lahoa karmissa tai puitteissa\n2️⃣ Huurretta lasien välissä\n3️⃣ Vetoa, vaikka tiivisteet on vaihdettu\n\nJos ongelma on pelkkä veto ehjissä ikkunoissa, huolto riittää usein. Sanomme sen myös suoraan.',
     'ig-3-merkit.png (somepaketti)', 'ikkunakauppias.fi/kannattaako-ikkunoiden-vaihto/', 0),
    ('2026-11-14', TT, 'Video 10–15 s', 'Osallistu', 'Arvaa mikä?',
     'Arvaa mikä tämä on? 👀 Vastaus videon lopussa. #ikkunat #arvaa',
     'Erikoislähikuva ikkunan painikkeesta tai saranasta, lopussa veto taaksepäin ja paljastus', '–', 0),
    # --- Vko 47: kotimaisuus ---
    ('2026-11-16', FB_IG, 'Karuselli (3 kuvaa)', 'Kotimaisuus', 'Mistä ikkunasi tulevat? Ylihärmästä',
     'Mistä uudet ikkunasi tulevat? 🇫🇮\n\nSkaala-ikkunat ja -ovet valmistetaan mittatilaustyönä Ylihärmässä. Jokainen ikkuna tehdään juuri sinun aukkoosi, eikä vakiokokoa tarvitse sovittaa vanhaan taloon.\n\nKotimaisuus & ammattitaito kulkee käsi kädessä.',
     'Tuotekuvat: olohuone-puuikkunat, ikkunan-rakenne, logo', 'ikkunakauppias.fi/ikkunat/', 15),
    ('2026-11-18', VIDEO, 'Video 30 s', 'Kulissit', 'Mitä mittauskäynnillä tapahtuu?',
     'Mitä ilmaisella mittauskäynnillä tapahtuu? 📏 Mittaamme aukot, katsomme ikkunoiden kunnon ja käymme läpi toiveesi. Käynti kestää yleensä alle tunnin, eikä se sido mihinkään.',
     'Kuvaa omaa mittausta (kysy asiakkaalta lupa, ei osoitteita näkyviin)', 'Varaa mittaus: linkki profiilissa', 0),
    ('2026-11-20', FB_IG, 'Kuva + äänestys tarinassa', 'Osallistu', 'Sisään- vai ulosaukeava?',
     'Kumpi on sinun suosikkisi? 🗳️\n\nA) Sisäänaukeava – pesu onnistuu sisältä\nB) Ulosaukeava – ikkunalauta jää vapaaksi\n\nVastaa kommentteihin A tai B!',
     'Kaksi kuvaa vierekkäin (ikkuna-verhot ja olohuone-puuikkunat)', 'Kommentoi A/B', 0),
    ('2026-11-21', TT, 'Video 15 s', 'Vinkki', 'Ikkunan rakenne 15 sekunnissa',
     'Mistä ikkuna koostuu? Karmi, puite, lasit ja tiivisteet 15 sekunnissa. #ikkunat #tiesitkö',
     'Ikkunan-rakenne-kuva tai oikea ikkuna, osoita sormella osat', '–', 0),
    # --- Vko 48: ovet, Black Friday ilman hätiköintiä ---
    ('2026-11-23', FB_IG, 'Kuva', 'Tuote', 'Ulko-ovi on talon käyntikortti',
     'Ulko-ovi on talon käyntikortti 🚪\n\nVanha ovi vetää, jumittaa tai huurtuu talvella? Skaalan mallistoista löytyy ovi sekä rintamamiestaloon että moderniin kotiin. Ulko-oven vaihto tehdään yleensä yhdessä päivässä.',
     'ig-4-ulko-ovi.png (somepaketti)', 'ikkunakauppias.fi/ovet/', 0),
    ('2026-11-25', VIDEO, 'Video 20 s', 'Vinkki', 'Ulko-oven 3 vaihdon merkkiä',
     'Pitäisikö ulko-ovi vaihtaa? 3 merkkiä: 1) vetää säätämisen jälkeenkin, 2) huurre tai jää sisäpinnalla, 3) lahoa alaosassa tai kynnyksessä. 🚪❄️',
     'Näytä vanhaa ovea: alaosa, lukko, karmin reuna', 'Linkki profiilissa', 0),
    ('2026-11-27', FB_IG, 'Kuva', 'Arvot', 'Black Friday – ilman kiirettä',
     'Black Friday? Meillä ei ole "vain tänään" -hintoja. 🙂\n\nMittauskäynti on ilmainen ja tarjous kirjallinen. Saat miettiä rauhassa, ja tarjouksesta näet, mitä hintaan sisältyy. Ikkunat ovat 40–60 vuoden hankinta, sitä ei kannata tehdä kiireessä.',
     'Rauhallinen kuva: ikkuna ja kahvikuppi, teksti "Ei kiirettä."', 'ikkunakauppias.fi/yhteystiedot/', 20),
    ('2026-11-28', TT, 'Video 20 s', 'Vinkki', 'Mitä ikkunaremontti maksaa?',
     'Mitä ikkunan vaihto maksaa? Yksi ikkuna asennettuna tyypillisesti 400–1 500 €, omakotitalon ikkunat yleensä 6 000–20 000 €. Tarkka hinta selviää ilmaisella mittauksella. 💶',
     'Sami kameralle tai teksti näytölle, taustalla ikkuna', 'Linkki profiilissa', 0),
    # --- Vko 49: talviasennus, itsenäisyyspäivä ---
    ('2026-11-30', FB_IG, 'Kuva', 'Vinkki', 'Voiko ikkunat vaihtaa talvella? Voi.',
     'Voiko ikkunat vaihtaa talvella? Kyllä voi. ❄️\n\nIkkunat vaihdetaan yksi kerrallaan, ja aukko on auki vain hetken. Koti ei ehdi jäähtyä, ja talvella aikataulut ovat usein joustavampia kuin keväällä.',
     'puuikkuna-kaihtimet (talvikuva)', 'ikkunakauppias.fi/ikkunaremontti/', 0),
    ('2026-12-02', VIDEO, 'Video 30 s, timelapse', 'Kulissit', 'Ikkunan vaihto pakkasella',
     'Ikkunan vaihto pakkasella alusta loppuun ⏱️ Vanha ulos, aukko suojaan, uusi paikalleen, tiivistys. Koti ei ehdi jäähtyä.',
     'Timelapse yhdestä vaihdosta (puhelin jalustalle), kansi tiktok-3-asennus.png', 'Linkki profiilissa', 0),
    ('2026-12-04', FB_IG, 'Karuselli (3 kuvaa)', 'Kotimaisuus', 'Suomessa tehty',
     'Suomessa tehty. 🇫🇮\n\nKun valitset Skaala-ikkunat, ne valmistetaan Suomessa ja asentaa ammattilainen, joka tuntee suomalaiset talot ja talvet. Siksi kotimaisuus ei meille ole iskulause vaan tapa tehdä työtä.',
     'Logo + kotimaiset tuotekuvat', 'ikkunakauppias.fi', 15),
    ('2026-12-06', KAIKKI, 'Kuva', 'Kausi', 'Hyvää itsenäisyyspäivää',
     'Hyvää itsenäisyyspäivää! 🇫🇮🕯️🕯️\n\nIllalla klo 18 ikkunoille syttyy kaksi kynttilää. Kauniita ikkunoita ja lämpimiä koteja koko Suomeen.',
     'Kaksi kynttilää ikkunalaudalla, sinivalkoinen sävy', '–', 0),
    # --- Vko 50: energia ja laskuri ---
    ('2026-12-07', FB_IG, 'Kuva', 'Vinkki', 'Paljonko vanhat ikkunasi maksavat joka talvi?',
     'Paljonko vanhat ikkunasi maksavat joka talvi? 💸\n\nKokeile sivuiltamme energiansäästölaskuria: valitse ikkunoiden määrä, ikä ja lämmitystapa, ja näet arvion säästöstä euroina vuodessa.',
     'Kuvakaappaus laskurista + logo', 'ikkunakauppias.fi/energiansaastolaskuri/', 15),
    ('2026-12-09', VIDEO, 'Video 20 s', 'Vinkki', 'Laskuri puhelimella',
     'Laskin, paljonko omakotitalo säästäisi uusilla ikkunoilla 👇 Kokeile itse, linkki profiilissa. #energiansäästö',
     'Näyttötallenne laskurin käytöstä puhelimella', 'Linkki profiilissa', 0),
    ('2026-12-11', FB_IG, 'Karuselli (3 kuvaa)', 'Vinkki', 'Kolmi- vai nelilasinen?',
     'Kolmi- vai nelilasinen ikkuna? 🤔\n\nKolmilasinen riittää useimpiin taloihin. Nelilasinen sopii erityisesti matalaenergiataloihin ja meluisille paikoille. Kerromme mittauskäynnillä, kumpi sopii teidän taloonne.',
     'ikkunan-rakenne + 2 tekstikuvaa', 'ikkunakauppias.fi/kolmilasinen-vai-nelilasinen-ikkuna/', 0),
    ('2026-12-12', TT, 'Video 20 s', 'Vinkki', 'U-arvo 20 sekunnissa',
     'Mikä ihmeen U-arvo? Mitä pienempi luku, sitä vähemmän lämpöä karkaa. Vanha 2-lasinen ikkuna n. 2,5–2,8, nykyaikainen 0,6–1,0. 🌡️',
     'Sami kameralle tai teksti näytölle', '–', 0),
    # --- Vko 51: kotitalousvähennys, asennuspäivä ---
    ('2026-12-14', FB_IG, 'Kuva', 'Vinkki', 'Kotitalousvähennys ja vuodenvaihde',
     'Muistathan: asennustyöstä saa kotitalousvähennystä 🧾\n\nVähennys tehdään sen vuoden verotuksessa, jolloin työ maksetaan. Tarkista ajantasaiset ehdot Verohallinnon sivuilta, ja kysy meiltä, paljonko tarjouksessa on työn osuutta.',
     'Selkeä tekstikuva somepaketin tyylillä', 'ikkunakauppias.fi/kotitalousvahennys-ikkuna-ovi/', 0),
    ('2026-12-16', VIDEO, 'Video 30 s', 'Kulissit', 'Mitä asennuspäivänä tapahtuu?',
     'Mitä asennuspäivänä tapahtuu? 🛠️ Suojaamme lattiat, puramme vanhat, asennamme ja tiivistämme uudet ja siivoamme jälkemme. Vanhat ikkunat viemme kierrätykseen.',
     'Lyhyitä klippejä työmaalta (lupa asiakkaalta)', 'Linkki profiilissa', 0),
    ('2026-12-18', FB_IG, 'Kuva + kysely tarinassa', 'Osallistu', 'Mikä ikkunoissa ärsyttää eniten?',
     'Mikä teidän ikkunoissa ärsyttää eniten? 😅\n\n🌬️ Veto\n❄️ Huurre\n🖌️ Maalaaminen\n🔊 Melu\n\nKerro kommenteissa, vastaamme parhaisiin vinkeillä tammikuussa.',
     'Neljä ikonia tai emoji-grafiikka', 'Kommentoi', 0),
    ('2026-12-19', TT, 'Video 15 s', 'Kulissit', 'Vanhan ikkunan purku',
     'Vanha ikkuna ulos 15 sekunnissa 💪 #ikkunaremontti #satisfying',
     'Purkuklippi työmaalta', '–', 0),
    # --- Vko 52: joulu ---
    ('2026-12-21', FB_IG, 'Kuva', 'Kausi', 'Joulun valot ikkunoissa',
     'Joulun valot näkyvät kauneimmin ikkunoista. ✨\n\nMikä on teidän ikkunalaudan jouluperinne? Kynttilät, tähti vai tontut?',
     'ikkuna-verhot.webp tai oma jouluinen ikkunakuva', 'Kommentoi', 0),
    ('2026-12-23', KAIKKI, 'Kuva', 'Kausi', 'Hyvää joulua',
     'Rauhallista joulua ja kiitos ensimmäisistä viikoista! 🎄 Vastaamme viesteihin taas [tarkista päivä].',
     'Logo-kuva jouluisella sävyllä', '–', 0),
    # --- Vko 53: vuodenvaihde ---
    ('2026-12-28', FB_IG, 'Karuselli (4 kuvaa)', 'Vinkki', 'Kysytyimmät kysymykset',
     'Kysytyimmät kysymykset tähän mennessä 👉\n\n1. Mitä ikkunaremontti maksaa?\n2. Voiko ikkunat vaihtaa talvella?\n3. Kauanko vaihto kestää?\n4. Saako kotitalousvähennystä?\n\nVastaukset kuvissa. Mitä haluaisit kysyä?',
     '4 tekstikuvaa (kysymys + lyhyt vastaus)', 'ikkunakauppias.fi', 0),
    ('2026-12-30', VIDEO, 'Video 20 s', 'Tarjous', 'Kevään remontti: nyt on hyvä aika varata mittaus',
     'Suunnitteletko ikkunaremonttia keväälle? 🌱 Nyt on hyvä aika varata mittaus: tarjous ehtii rauhassa, ja asennusaika sovitaan ennen kevään ruuhkaa.',
     'Sami kameralle, lumimaisema taustalla', 'Varaa mittaus: linkki profiilissa', 15),
    ('2027-01-01', FB_IG, 'Kuva', 'Kausi', 'Hyvää uutta vuotta',
     'Hyvää uutta vuotta 2027! 🎆 Kiitos, että olet seurannut matkaamme.',
     'Logo-kuva', '–', 0),
    ('2027-01-02', TT, 'Video 20 s', 'Vinkki', '3 asiaa, jotka tarkistaa ikkunoista tammikuussa',
     '3 asiaa, jotka kannattaa tarkistaa ikkunoista tammikuussa ❄️ 1) huurre lasien välissä 2) veto reunoilla 3) jää karmin nurkissa. #ikkunat #talvi',
     'Lyhyet klipit omasta ikkunasta', '–', 0),
    # --- Vko 1: pakkaset ja lämmityslasku ---
    ('2027-01-04', FB_IG, 'Kuva', 'Vinkki', 'Tammikuun lämmityslasku',
     'Tammikuun lämmityslasku saapui? 🥶\n\nVanhojen ikkunoiden läpi karkaa moninkertaisesti lämpöä uusiin verrattuna. Laske oman talosi arvio energialaskurilla ja lue, milloin vaihto kannattaa ja milloin ei.',
     'olohuone-maisemaikkuna + tekstipalkki', 'ikkunakauppias.fi/kannattaako-ikkunoiden-vaihto/', 20),
    ('2027-01-06', VIDEO, 'Video 20 s', 'Vinkki', 'Pakkastesti: missä ikkuna on kylmä?',
     'Pakkastesti 🥶 Tunnustele kämmenselällä puitteen reunat ja lasin alaosa. Kylmä kohta kertoo vuodosta tai väsyneestä tiivisteestä.',
     'Käsi ikkunan reunoilla, ulkona pakkasta', 'Linkki profiilissa', 0),
    ('2027-01-08', FB_IG, 'Karuselli (3 kuvaa)', 'Vinkki', 'Vaihtaa vai huoltaa?',
     'Uudet ikkunat vai huolto? 🤔\n\n✅ Huolto riittää: veto, mutta puu ehjää ja puitteet suorat.\n🔁 Vaihto kannattaa: lahoa karmissa, huurretta lasien välissä, ikkunat yli 40 v.\n\nSanomme mittauskäynnillä suoraan, jos ikkunasi eivät vielä tarvitse vaihtoa.',
     '3 tekstikuvaa somepaketin tyylillä', 'ikkunakauppias.fi/ikkunaremontti-vai-kunnostus/', 0),
    ('2027-01-09', TT, 'Video 25 s', 'Arvot', 'Maksavatko uudet ikkunat itsensä takaisin?',
     'Maksavatko uudet ikkunat itsensä takaisin parissa vuodessa? Rehellinen vastaus: eivät pelkällä energiansäästöllä. Vaihto kannattaa, kun ikkunat ovat käyttöikänsä lopussa. 👇',
     'Sami kameralle', 'Linkki profiilissa', 0),
    # --- Vko 2: ovet ---
    ('2027-01-11', FB_IG, 'Kuva', 'Tuote', 'Skaala Aika -ovi',
     'Tumma ulko-ovi, joka kestää Suomen talvet? 🖤\n\nSkaala Aika on alumiinipintainen ovi, joka kestää myös tummat sävyt. Pinnalle 20 vuoden takuu.',
     'Aika-oven tuotekuva (pyydä Skaalalta) tai ulko-ovi-terassi', 'ikkunakauppias.fi/skaala-aika-ovi/', 15),
    ('2027-01-13', VIDEO, 'Video 20 s', 'Vinkki', 'Jäätyykö ulko-ovi?',
     'Jäätyykö ulko-ovesi lukko tai karmin nurkka? ❄️ Se kertoo, että ovi johtaa kylmää sisään tai ei ole tiivis. Tiivisteet voi vaihtaa, mutta vanhassa ovessa vika ei yleensä poistu huoltamalla.',
     'Lähikuva jäisestä lukosta tai karmista', 'Linkki profiilissa', 0),
    ('2027-01-15', FB_IG, 'Kuva + äänestys tarinassa', 'Osallistu', 'Kumpi ovi?',
     'Kumpi ulko-ovi sopisi teidän taloonne? 🚪\n\nA) Tumma ja moderni\nB) Perinteinen ristiuritettu\n\nVastaa A tai B!',
     'Kaksi ovikuvaa vierekkäin (Skaalan kuvapankki)', 'Kommentoi A/B', 0),
    ('2027-01-16', TT, 'Video 30 s, timelapse', 'Kulissit', 'Ulko-oven vaihto yhdessä päivässä',
     'Ulko-oven vaihto yhdessä päivässä ⏱️🚪 #oviremontti #ennenjajälkeen',
     'Timelapse oven vaihdosta, lopussa valmis ovi', '–', 0),
    # --- Vko 3: kevään suunnittelu, taloyhtiöt ---
    ('2027-01-18', FB_IG, 'Kuva', 'Tarjous', 'Kevään remontti alkaa tammikuussa',
     'Kevään ikkunaremontti alkaa tammikuussa 📅\n\nMittaus → tarjous → valmistus Suomessa → asennus. Kun mittaus tehdään nyt, asennusaika sovitaan rauhassa ennen kevään ruuhkaa.',
     'Aikajanagrafiikka somepaketin tyylillä', 'ikkunakauppias.fi/yhteystiedot/', 20),
    ('2027-01-20', VIDEO, 'Video 30–45 s', 'Osallistu', 'Vastaamme kysymyksiinne',
     'Vastaamme kysymyksiinne! 🙋 Tällä kertaa: [valitse 2–3 kysymystä kommenteista ja tarinoista].',
     'Sami kameralle, kysymykset tekstinä näytölle', 'Kysy lisää kommenteissa', 0),
    ('2027-01-22', FB_IG, 'Karuselli (4 kuvaa)', 'Vinkki', 'Taloyhtiön ikkunaremontti 4 vaiheessa',
     'Taloyhtiön ikkunaremontti 4 vaiheessa 🏢\n\n1. Kuntoarvio ja tarve\n2. Tarjoukset hallitukselle\n3. Päätös yhtiökokouksessa\n4. Asennus porras kerrallaan\n\nVoimme tulla kertomaan vaihtoehdoista myös hallituksen kokoukseen.',
     '4 tekstikuvaa', 'ikkunakauppias.fi/taloyhtion-ikkunaremontti/', 0),
    ('2027-01-23', TT, 'Video 15 s', 'Vinkki', 'Mökin ikkunat talvella',
     'Mökki kylmillään talvella? Tarkista ikkunat keväällä: huurre ja kosteus lasien välissä kertovat vuodosta. 🏡❄️ #mökki',
     'Mökkikuva tai lumimaisema', '–', 0),
    # --- Vko 4: yhteenveto ja CTA ---
    ('2027-01-25', FB_IG, 'Kuva', 'Tarjous', 'Ilmainen mittauskäynti',
     'Ilmainen mittauskäynti 📏\n\nTulemme paikan päälle, mittaamme aukot ja annamme kirjallisen tarjouksen. Ei kiirettä, ei sitoumuksia.',
     'ig-2-mittauskaynti.png (somepaketti)', 'ikkunakauppias.fi/yhteystiedot/', 20),
    ('2027-01-27', VIDEO, 'Video 20–30 s', 'Tutuksi', 'Kiitos ensimmäisistä kuukausista',
     'Kiitos, että olette seuranneet! 🙏 Kevään aikana luvassa lisää vinkkejä, työmaakuvia ja vastauksia kysymyksiinne. Mitä haluaisit nähdä?',
     'Sami kameralle', 'Kommentoi toiveesi', 0),
    ('2027-01-29', FB_IG, 'Kuva', 'Osallistu', 'Kysy meiltä',
     'Kysy meiltä mitä vain ikkunoista ja ovista 💬\n\nVastaamme viesteihin ja kommentteihin, ja sivuiltamme löytyy Kysy meiltä -ikkuna yleisimpiin kysymyksiin.',
     'Logo + chat-ikoni', 'ikkunakauppias.fi', 0),
    ('2027-01-30', TT, 'Video 15 s', 'Vinkki', 'Talvivinkki: ikkunan tiivisteet',
     'Talvivinkki: pyyhi tiivisteet ja voitele saranat kerran vuodessa, niin ikkuna sulkeutuu tiiviisti. 🧽 #remonttivinkit',
     'Lähikuva tiivisteestä ja saranasta', '–', 0),
]

VKPAIVA = ['Ma', 'Ti', 'Ke', 'To', 'Pe', 'La', 'Su']

def build(path):
    wb = Workbook()
    font = 'Arial'
    head_fill = PatternFill('solid', fgColor='16303B')
    gold = PatternFill('solid', fgColor='FFD866')
    week_fill = PatternFill('solid', fgColor='EBF1F3')
    thin = Side(style='thin', color='D5DEE2')
    border = Border(bottom=thin)

    # ---------------- Ohje ----------------
    oh = wb.active
    oh.title = 'Ohje ja tavoitteet'
    oh.sheet_view.showGridLines = False
    oh.column_dimensions['A'].width = 34
    oh.column_dimensions['B'].width = 70
    rows = [
        ('SOMEKALENTERI marraskuu 2026 – tammikuu 2027', None),
        ('Ikkunakauppias.fi – Nordic Ikkunat & Ovet Oy', None),
        (None, None),
        ('Tavoite', 'Tunnettuuden kasvattaminen: seuraajat, kattavuus ja tutuksi tuleminen Turun seudulla ja Varsinais-Suomessa.'),
        ('Mittarit', 'Seuraajamäärä (FB, IG, TikTok), kattavuus/näyttökerrat viikossa, kommentit ja jaot. Kirjaa luvut Seuranta-välilehdelle maanantaisin.'),
        ('Rytmi', 'Ma: kuva tai karuselli (FB + IG) · Ke: video (Reels + TikTok) · Pe: kuva, karuselli tai kysely (FB + IG) · La: lyhyt TikTok. Lisäksi juhlapäivät.'),
        ('Tarinat (stories)', '2–4 kertaa viikossa, kevyesti: työmaan arkea, kysely- ja kysymystarrat, jaa uusi postaus tarinaan. Ei tarvitse olla hiottua.'),
        ('Julkaisuajat', 'Arkisin klo 18–20 tai aamulla klo 7–8. Viikonloppuna klo 10–12. Seuraa myöhemmin omista tilastoista, mikä toimii.'),
        ('Tehostus', 'Merkityt julkaisut kannattaa tehostaa pienellä summalla (15–20 €), kohderyhmä 30–70-vuotiaat, Turku + 40 km. Facebookissa ja Instagramissa sama mainos.'),
        ('Vinkit', 'Kasvot ja oma ääni toimivat paremmin kuin mainoskuvat. Vastaa jokaiseen kommenttiin samana päivänä. Pyydä asiakkailta lupa työmaakuviin, älä näytä osoitteita.'),
        ('Muokattavat sarakkeet', 'Kalenteri-välilehdellä Tila (valikko) ja Muistiinpanot. Muut sarakkeet ovat valmis suunnitelma, joita voi muokata vapaasti.'),
        (None, None),
        ('Yhteenveto', None),
        ('Julkaisuja yhteensä', "=COUNTA(Kalenteri!D2:D%d)" % (len(P) + 1 + 13)),
        ('Julkaistu', "=COUNTIF(Kalenteri!L:L,\"Julkaistu\")"),
        ('Valmiina odottamassa', "=COUNTIF(Kalenteri!L:L,\"Valmis\")"),
        ('Tehostusbudjetti yhteensä (€)', "=SUM(Kalenteri!K:K)"),
        (None, None),
        ('Julkaisut teemoittain', None),
    ]
    teemat = ['Tutuksi', 'Vinkki', 'Kotimaisuus', 'Kulissit', 'Osallistu', 'Tuote', 'Arvot', 'Kausi', 'Tarjous']
    for t in teemat:
        rows.append((t, '=COUNTIF(Kalenteri!F:F,"%s")' % t))
    for i, (a, b) in enumerate(rows, start=1):
        if a is not None:
            oh.cell(i, 1, a)
        if b is not None:
            oh.cell(i, 2, b)
        for c in (1, 2):
            cell = oh.cell(i, c)
            cell.font = Font(name=font, size=11, bold=(c == 1))
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    oh['A1'].font = Font(name=font, size=16, bold=True, color='16303B')
    oh['A2'].font = Font(name=font, size=11, color='647A84')
    for r in (13, 19):
        oh.cell(r, 1).font = Font(name=font, size=12, bold=True, color='1D5E79')
    for r in range(14, 18):
        oh.cell(r, 2).alignment = Alignment(horizontal='left')
    oh['B17'].number_format = '0 "€"'
    for r in range(2, len(rows) + 1):
        if oh.cell(r, 2).value and isinstance(oh.cell(r, 2).value, str) and not oh.cell(r, 2).value.startswith('='):
            oh.row_dimensions[r].height = 32

    # ---------------- Kalenteri ----------------
    ws = wb.create_sheet('Kalenteri')
    headers = ['Viikko', 'Päivä', 'Pvm', 'Kanavat', 'Muoto', 'Teema', 'Aihe', 'Julkaisuteksti (kopioi sellaisenaan)',
               'Kuva / video', 'Linkki / CTA', 'Tehostus €', 'Tila', 'Muistiinpanot']
    widths = [8, 7, 11, 22, 18, 12, 30, 70, 34, 30, 11, 13, 24]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(1, i, h)
        c.font = Font(name=font, bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'D2'

    r = 2
    last_week = None
    for (d, kan, muoto, teema, aihe, teksti, kuva, cta, teho) in P:
        date = dt.date.fromisoformat(d)
        wk = date.isocalendar()[1]
        if wk != last_week:
            ws.cell(r, 1, 'Vko %d' % wk)
            ws.cell(r, 2, 'Viikon tarinat: 2–4 kpl (työmaa, kysely, jaa postaus)')
            for c in range(1, 14):
                ws.cell(r, c).fill = week_fill
                ws.cell(r, c).font = Font(name=font, bold=True, color='1D5E79', size=10)
            r += 1
            last_week = wk
        caption = teksti + ('\n\n' + H if kan != TT else '')
        vals = [wk, VKPAIVA[date.weekday()], date, kan, muoto, teema, aihe, caption, kuva, cta, teho or None, 'Suunniteltu', None]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=font, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = border
        ws.cell(r, 3).number_format = 'd.m.yyyy'
        ws.cell(r, 11).number_format = '0 "€"'
        ws.cell(r, 7).font = Font(name=font, size=10, bold=True)
        lines = caption.count('\n') + len(caption) // 85 + 1
        ws.row_dimensions[r].height = max(45, min(260, 13 * lines))
        r += 1
    last = r - 1
    dv = DataValidation(type='list', formula1='"Suunniteltu,Valmis,Julkaistu,Siirretty"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('L2:L%d' % last)
    ws.conditional_formatting.add('L2:L%d' % last, FormulaRule(formula=['$L2="Julkaistu"'], fill=PatternFill('solid', fgColor='D7F0DE')))
    ws.conditional_formatting.add('L2:L%d' % last, FormulaRule(formula=['$L2="Valmis"'], fill=gold))
    ws.auto_filter.ref = 'A1:M%d' % last
    oh['B14'] = '=COUNTIFS(Kalenteri!G2:G%d,"<>")' % last

    # ---------------- Seuranta ----------------
    se = wb.create_sheet('Seuranta')
    sh = ['Viikko', 'Maanantai', 'Facebook-seuraajat', 'Instagram-seuraajat', 'TikTok-seuraajat', 'Yhteensä',
          'Kasvu edellisestä', 'Kattavuus viikossa (kaikki)', 'Kommentit + jaot']
    for i, h in enumerate(sh, start=1):
        c = se.cell(1, i, h)
        c.font = Font(name=font, bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
        se.column_dimensions[get_column_letter(i)].width = 14 if i > 2 else 10
    se.column_dimensions['B'].width = 12
    se.row_dimensions[1].height = 34
    start = dt.date(2026, 11, 2)
    for i in range(14):
        row = i + 2
        mon = start + dt.timedelta(weeks=i)
        se.cell(row, 1, 'Vko %d' % mon.isocalendar()[1])
        se.cell(row, 2, mon).number_format = 'd.m.yyyy'
        for c in (3, 4, 5, 8, 9):
            se.cell(row, c).font = Font(name=font, color='0000FF')
            se.cell(row, c).fill = PatternFill('solid', fgColor='FFF7D6')
        se.cell(row, 6, '=IF(COUNT(C%d:E%d)=0,"",SUM(C%d:E%d))' % (row, row, row, row))
        se.cell(row, 7, '' if row == 2 else '=IF(OR(F%d="",F%d=""),"",F%d-F%d)' % (row, row - 1, row, row - 1))
        for c in range(1, 10):
            se.cell(row, c).font = Font(name=font, size=10, color='0000FF' if c in (3, 4, 5, 8, 9) else '000000')
    se['A17'] = 'Täytä keltaiset solut maanantaisin (luvut sovellusten tilastoista). Yhteensä ja kasvu lasketaan automaattisesti.'
    se['A17'].font = Font(name=font, size=10, italic=True, color='647A84')
    se['A18'] = 'Esimerkki: vko 45 FB 120, IG 85, TikTok 40 → yhteensä 245.'
    se['A18'].font = Font(name=font, size=10, italic=True, color='647A84')

    wb.calculation.fullCalcOnLoad = True  # laskee kaavat avattaessa
    wb.save(path)


if __name__ == '__main__':
    build(sys.argv[1])
